"""Read and validate source data without changing the Excel workbook."""

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from math import isfinite
from numbers import Real
from pathlib import Path
from types import MappingProxyType
from typing import Mapping
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


SOURCE_SHA256 = "ba2ae447d1964e5bd0d2d92d89d98b837be28b40a81f9ad1720b5078391eccbf"
HEADERS = {
    "ROUTES": ("Jeep Code", "Stop 1", "Stop 2", "Stop 3", "Stop 4", "Stop 5"),
    "NODE COSTS": (
        "Place", "Congestion index (0-10)", "Transfer penalty (min)",
        "Safety risk (0-10)", "Access penalty (min)", "Note",
    ),
    "EDGE COSTS": (
        "Jeep Code", "From", "To", "Travel time (min)",
        "Incremental fare (PHP)", "Risk (0-10)", "Congestion multiplier", "Status",
    ),
}


class ValidationError(ValueError):
    """An actionable workbook or lookup error suitable for display to a user."""


@dataclass(frozen=True)
class Service:
    code: str
    stops: tuple[str, ...]
    source_row: int


@dataclass(frozen=True)
class Node:
    place: str
    congestion_index: float
    transfer_minutes: float
    safety_risk: float
    access_minutes: float
    note: str
    source_row: int


@dataclass(frozen=True)
class Segment:
    code: str
    origin: str
    destination: str
    travel_minutes: float
    incremental_fare: float
    risk: float
    congestion_multiplier: float
    status: str
    source_row: int
    disruption_multiplier: float = 1.0
    driving_minutes_override: float | None = None

    @property
    def driving_minutes(self) -> float:
        if self.driving_minutes_override is not None:
            return _number(self.driving_minutes_override, "Driving-time override")
        return _number(self.travel_minutes * self.congestion_multiplier * self.disruption_multiplier,
                       "Calculated driving minutes")

    @property
    def source_id(self) -> tuple[str, int]:
        return "EDGE COSTS", self.source_row


@dataclass(frozen=True)
class Network:
    services: Mapping[str, Service]
    nodes: Mapping[str, Node]
    segments: tuple[Segment, ...]
    adjacency: Mapping[str, tuple[Segment, ...]]
    source_path: Path | None = None
    source_sha256: str | None = None
    max_transfers: int | None = None
    scenario_id: str = "S1"

    def service(self, code: str) -> Service:
        if code not in self.services:
            raise ValidationError(f"Unknown jeepney code: {code!r}. Choose a listed code.")
        return self.services[code]

    def require_stop(self, place: str) -> Node:
        if place not in self.nodes:
            raise ValidationError(f"Unknown stop: {place!r}. Choose a listed stop.")
        return self.nodes[place]

    def outgoing(self, place: str) -> tuple[Segment, ...]:
        """Only OPEN segments may be used for travel, in either direction."""
        self.require_stop(place)
        return tuple(edge for edge in self.adjacency[place] if edge.status == "OPEN")


def _text(value, where: str) -> str:
    if not isinstance(value, str) or not value or value != value.strip():
        raise ValidationError(f"{where}: expected nonempty text without surrounding spaces.")
    if value.startswith("="):
        raise ValidationError(f"{where}: formulas are not supported in source data.")
    return value


def _number(value, where: str, maximum: float | None = None) -> float:
    if isinstance(value, bool) or not isinstance(value, Real):
        raise ValidationError(f"{where}: expected a finite nonnegative number.")
    try:
        number = float(value)
    except OverflowError as exc:
        raise ValidationError(f"{where}: number exceeds the supported finite numeric range.") from exc
    if not isfinite(number) or number < 0 or (maximum is not None and number > maximum):
        limit = f" from 0 to {maximum:g}" if maximum is not None else " >= 0"
        raise ValidationError(f"{where}: expected a finite number{limit}.")
    return number


def _records(sheet: str, rows):
    """Resolve fields by exact header name, including reordered named columns."""
    rows = iter(rows)
    header = list(next(rows, ()))
    while header and header[-1] is None:
        header.pop()
    expected = HEADERS[sheet]
    if (len(header) != len(expected) or any(not isinstance(h, str) for h in header)
            or set(header) != set(expected)):
        raise ValidationError(f"{sheet} row 1: expected exactly these headers: {', '.join(expected)}.")
    blank_row = None
    for row_number, values in enumerate(rows, start=2):
        values = tuple(values)
        if all(value is None for value in values):
            if blank_row is None:
                blank_row = row_number
            continue
        if blank_row is not None:
            raise ValidationError(f"{sheet} row {blank_row}: empty rows are only allowed at the end of a table.")
        if any(value is not None for value in values[len(header):]):
            raise ValidationError(f"{sheet} row {row_number}: data appears in an unnamed column.")
        padded = values + (None,) * max(0, len(header) - len(values))
        yield row_number, dict(zip(header, padded))


def validate_tables(tables: Mapping[str, object]) -> Network:
    """Validate raw worksheet rows and build an immutable service-aware graph."""
    for sheet in HEADERS:
        if sheet not in tables:
            raise ValidationError(f"Missing required worksheet: {sheet}.")

    services = {}
    for row, record in _records("ROUTES", tables["ROUTES"]):
        where = f"ROUTES row {row}"
        code = _text(record["Jeep Code"], f"{where}, Jeep Code")
        if code in services:
            raise ValidationError(f"{where}: duplicate jeepney code {code!r}.")
        stops = tuple(_text(record[f"Stop {i}"], f"{where}, Stop {i}") for i in range(1, 6))
        if len(set(stops)) != len(stops):
            raise ValidationError(f"{where}: repeated stops are not allowed in a service.")
        services[code] = Service(code, stops, row)

    nodes = {}
    for row, record in _records("NODE COSTS", tables["NODE COSTS"]):
        where = f"NODE COSTS row {row}"
        place = _text(record["Place"], f"{where}, Place")
        if place in nodes:
            raise ValidationError(f"{where}: duplicate stop {place!r}.")
        values = [
            _number(record[field], f"{where}, {field}", maximum)
            for field, maximum in zip(HEADERS["NODE COSTS"][1:5], (10, None, 10, None))
        ]
        note = record["Note"]
        if note is not None:
            note = _text(note, f"{where}, Note")
        nodes[place] = Node(place, *values, note or "", row)

    route_stops = {stop for service in services.values() for stop in service.stops}
    if route_stops != set(nodes):
        raise ValidationError(
            f"Stop membership mismatch: missing NODE COSTS {sorted(route_stops - set(nodes))}; "
            f"unlisted in ROUTES {sorted(set(nodes) - route_stops)}."
        )
    expected_pairs = {
        (service.code, frozenset((a, b)))
        for service in services.values()
        for a, b in zip(service.stops, service.stops[1:])
    }
    seen = set()
    segments = []
    adjacency = {place: [] for place in nodes}
    for row, record in _records("EDGE COSTS", tables["EDGE COSTS"]):
        where = f"EDGE COSTS row {row}"
        code, origin, destination = (
            _text(record[field], f"{where}, {field}") for field in HEADERS["EDGE COSTS"][:3]
        )
        key = code, frozenset((origin, destination))
        if key not in expected_pairs:
            raise ValidationError(f"{where}: {code} {origin} - {destination} is not a listed adjacent service pair.")
        if key in seen:
            raise ValidationError(f"{where}: duplicate service segment (including reverse duplicates).")
        seen.add(key)
        values = [
            _number(record[field], f"{where}, {field}", maximum)
            for field, maximum in zip(HEADERS["EDGE COSTS"][3:7], (None, None, 10, None))
        ]
        status = record["Status"]
        if status not in ("OPEN", "CLOSED"):
            raise ValidationError(f"{where}, Status: expected OPEN or CLOSED.")
        edge = Segment(code, origin, destination, *values, status, row)
        reverse = Segment(code, destination, origin, *values, status, row)
        segments.append(edge)
        adjacency[origin].append(edge)
        adjacency[destination].append(reverse)
    if seen != expected_pairs:
        missing = sorted((code, sorted(pair)) for code, pair in expected_pairs - seen)
        raise ValidationError(f"EDGE COSTS: missing service segments: {missing}.")
    if not services or not nodes or not segments:
        raise ValidationError("The source network must contain services, stops, and segments.")
    return Network(
        MappingProxyType(services), MappingProxyType(nodes), tuple(segments),
        MappingProxyType({place: tuple(edges) for place, edges in adjacency.items()}),
    )


def load_network(path: str | Path, expected_sha256: str | None = None) -> Network:
    """Read literal values only; close the workbook even when validation fails."""
    path = Path(path).resolve()
    try:
        source_bytes = path.read_bytes()
        fingerprint = sha256(source_bytes).hexdigest()
        if expected_sha256 is not None and fingerprint.lower() != expected_sha256.lower():
            raise ValidationError("Workbook fingerprint differs from the reviewed source. Check the file before continuing.")
        workbook = load_workbook(BytesIO(source_bytes), read_only=True, data_only=False)
        try:
            tables = {
                sheet: list(workbook[sheet].values)
                for sheet in HEADERS if sheet in workbook.sheetnames
            }
        finally:
            workbook.close()
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as exc:
        if isinstance(exc, ValidationError):
            raise
        raise ValidationError(f"Cannot read workbook {path.name}: {exc}") from exc
    network = validate_tables(tables)
    return Network(
        network.services, network.nodes, network.segments, network.adjacency,
        path, fingerprint,
    )
